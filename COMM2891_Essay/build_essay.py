"""Build the COMM2891 Major Research Essay as a formatted .docx.

Formatting follows the brief: title page, Times New Roman 12, double-spaced
throughout, first-line indent on every body paragraph, page numbers, default
(1 inch) margins, word count listed before the reference list.

Content lives in essay_content.py so text can be swapped without touching
the formatting code. Usage: python3 build_essay.py [output_name.docx]
"""
import datetime
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

import essay_content as C

FONT = "Times New Roman"
SIZE = Pt(12)


def set_run_font(run, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = SIZE
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def add_rich_text(p, text):
    """Add text to paragraph p; *asterisk-wrapped* spans are italicised."""
    for i, chunk in enumerate(re.split(r"\*(.+?)\*", text)):
        if not chunk:
            continue
        set_run_font(p.add_run(chunk), italic=(i % 2 == 1))


def fmt(p, first_indent=None, hanging=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if hanging is not None:
        pf.left_indent = hanging
        pf.first_line_indent = -hanging


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    set_run_font(run)
    for tag, text in (("begin", None), ("instr", " PAGE "), ("separate", None), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        run._element.append(el)


def add_heading(doc, text, level=1, center=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    fmt(p, first_indent=Inches(0.5))
    add_rich_text(p, text)
    return p


def count_words(text):
    return len(re.findall(r"\S+", text.replace("*", "")))


def body_word_count():
    return sum(count_words(t) for _, paras in C.SECTIONS for t in paras)


def heading_word_count():
    return sum(count_words(h) for h, _ in C.SECTIONS if h and h != "Introduction")


def build(out_path):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for lvl in (1, 2):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name = FONT
        hs.font.size = SIZE
        hs.font.bold = True
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.paragraph_format.space_before = Pt(0)
        hs.paragraph_format.space_after = Pt(0)
        hs.paragraph_format.line_spacing = 2.0
        rpr = hs.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Inches(1))
    add_page_number_footer(sec)

    # ---- Title page ----
    for _ in range(3):
        fmt(doc.add_paragraph())
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(p.add_run(C.TITLE), bold=True)
    fmt(doc.add_paragraph())
    for line in C.TITLE_PAGE_LINES:
        p = doc.add_paragraph()
        fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run_font(p.add_run(line))

    # ---- Essay body ----
    p = doc.add_paragraph()
    fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.page_break_before = True
    set_run_font(p.add_run(C.TITLE), bold=True)
    for heading, paras in C.SECTIONS:
        if heading and heading != "Introduction":  # APA 7: the paper title heads the introduction
            add_heading(doc, heading, level=1, center=True)
        for t in paras:
            add_body(doc, t)

    # ---- Word count (before references, per brief) ----
    body = body_word_count()
    p = doc.add_paragraph()
    fmt(p)
    set_run_font(
        p.add_run(f"Word count: {body:,} (excluding title page, headings and reference list)"),
        bold=True,
    )

    # ---- References ----
    h = add_heading(doc, "References", level=1, center=True)
    h.paragraph_format.page_break_before = True
    for ref in sorted(C.REFERENCES, key=lambda r: r.replace("*", "").lower()):
        p = doc.add_paragraph()
        fmt(p, hanging=Inches(0.5))
        add_rich_text(p, ref)

    cp = doc.core_properties
    cp.title = C.TITLE
    cp.subject = "COMM2891 Asian Media and Communication"
    cp.author = cp.last_modified_by = ""
    cp.comments = ""
    cp.keywords = ""
    cp.created = cp.modified = datetime.datetime.now()
    doc.save(out_path)
    return body


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "COMM2891_Fukushima_Essay_v1.docx"
    body = build(out)
    print(f"Saved {out}")
    print(f"Body words: {body} | headings: {heading_word_count()} | total incl. headings: {body + heading_word_count()}")
    for h, paras in C.SECTIONS:
        print(f"  {h or '(intro)'}: {sum(count_words(t) for t in paras)}")
